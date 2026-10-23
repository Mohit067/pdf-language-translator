"""Upload -> translate (real local IndicTrans2) -> preview -> download, for every supported format.
Skipped automatically when the model cannot be loaded on this machine."""

import csv
import io
import re
from pathlib import Path

import docx
import openpyxl
import pymupdf
import pytest

import backend.main as main
from tests.conftest import (
    ASSETS_DIR, CSV_TEXT, MD_TEXT, TXT_TEXT, make_docx, make_pdf, make_xlsx, submit, wait_for_job,
)

pytestmark = pytest.mark.model

DEVANAGARI = re.compile(r"[ऀ-ॿ]")
TAMIL = re.compile(r"[஀-௿]")


def translate_ok(client, filename, data, source="English", target="Hindi"):
    r = submit(client, filename, data, source, target)
    assert r.status_code == 202, r.text
    status = wait_for_job(client, r.json()["job_id"])
    assert status["status"] == "completed", status.get("error")
    result = status["result"]
    download = client.get(result["download_url"])
    assert download.status_code == 200
    assert "attachment" in download.headers["content-disposition"]
    return result, download.content


def test_txt_preserves_structure_and_numbers(client, require_model):
    require_model("en-indic")
    result, content = translate_ok(client, "notes.txt", TXT_TEXT.encode("utf-8"))
    out = content.decode("utf-8")
    assert out.count("\n") == TXT_TEXT.count("\n")             # line structure kept
    assert "Total: 4,500.00" in out or "4,500.00" in out
    lines = out.split("\n")
    assert DEVANAGARI.search(lines[0]) and DEVANAGARI.search(lines[2])
    assert "Meeting" not in out and "Marketing will" not in out
    preview = client.get(result["translated_url"])
    assert preview.headers["content-type"].startswith("text/plain")
    assert "inline" in preview.headers["content-disposition"]


def test_pdf_keeps_layout_images_and_graphics(client, require_model):
    require_model("en-indic")
    result, content = translate_ok(client, "report.pdf", make_pdf(with_image=True))
    src = pymupdf.open(stream=make_pdf(with_image=True), filetype="pdf")
    out = pymupdf.open(stream=content, filetype="pdf")
    assert out.page_count == src.page_count
    page = out[0]
    assert len(page.get_images()) == 1                          # image kept
    assert page.get_drawings()                                  # vector rule kept
    text = page.get_text()
    assert "The company expanded" not in text                   # original text replaced
    assert "4,500.00" in text                                   # numeric-only text untouched
    assert all(DEVANAGARI.search(c["translated"]) for c in result["chunks"])
    assert client.get(result["translated_url"]).headers["content-type"] == "application/pdf"
    assert result["download_filename"] == "translated_report.pdf"


def test_docx_keeps_formatting_images_tables_and_header(client, require_model):
    require_model("en-indic")
    result, content = translate_ok(client, "report.docx", make_docx())
    d = docx.Document(io.BytesIO(content))
    assert len(d.inline_shapes) == 1                            # picture kept
    body_runs = [p for p in d.paragraphs if p.runs and DEVANAGARI.search(p.text)]
    assert any(p.runs[0].bold for p in body_runs)               # first-run formatting kept
    assert d.paragraphs[-1].alignment == 1                      # centered paragraph stays centered
    table = d.tables[0]
    assert len(table.rows[0]._tr.tc_lst) == 2                   # merge preserved (2 physical cells)
    assert DEVANAGARI.search(table.cell(0, 0).text)
    assert table.cell(1, 1).text == "12,500"
    assert DEVANAGARI.search(d.sections[0].header.paragraphs[0].text)
    texts = [c["original"] for c in result["chunks"]]
    assert texts.count("Customer feedback") == 1


def test_xlsx_keeps_formulas_numbers_and_sheet_names(client, require_model):
    require_model("en-indic")
    result, content = translate_ok(client, "sales.xlsx", make_xlsx())
    wb = openpyxl.load_workbook(io.BytesIO(content))
    assert wb.sheetnames == ["Sales", "Notes"]
    ws = wb["Sales"]
    assert ws["D2"].value == "=B2*C2" and ws["B4"].value == "=SUM(B2:B3)"
    assert ws["B2"].value == 10 and ws["C2"].value == 25.5
    assert ws["F1"].value == "INV-2026-001"
    assert DEVANAGARI.search(ws["A1"].value)
    assert ws["A2"].value == wb["Notes"]["A1"].value            # same source -> same translation
    assert result["pages"] == 2


def test_csv_keeps_shape_and_quoting(client, require_model):
    require_model("en-indic")
    _, content = translate_ok(client, "staff.csv", CSV_TEXT.encode("utf-8"))
    assert content.startswith(b"\xef\xbb\xbf")                  # BOM for Excel
    rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
    src_rows = list(csv.reader(io.StringIO(CSV_TEXT)))
    assert [len(r) for r in rows] == [len(r) for r in src_rows]
    assert rows[3] == ["101", "202", "303"]
    assert DEVANAGARI.search(rows[1][2])


def test_markdown_keeps_syntax_and_code(client, require_model):
    require_model("en-indic")
    _, content = translate_ok(client, "README.md", MD_TEXT.encode("utf-8"))
    out = content.decode("utf-8")
    assert 'print("do not translate")' in out
    lines = out.split("\n")
    assert lines[0].startswith("# ") and DEVANAGARI.search(lines[0])
    assert lines[4].startswith("- ") and lines[5].startswith("- ")
    assert lines[11].count("|") == 3 and lines[12] == "|---------|--------|"


def test_english_to_tamil_outputs_tamil_script(client, require_model):
    require_model("en-indic")
    result, content = translate_ok(client, "t.txt", b"The meeting is tomorrow morning.", "English", "Tamil")
    out = content.decode("utf-8")
    assert TAMIL.search(out) and not DEVANAGARI.search(out)    # script restored from Devanagari


def test_hindi_to_english(client, require_model):
    require_model("indic-en")
    _, content = translate_ok(client, "h.txt", "आज मौसम बहुत अच्छा है।".encode("utf-8"), "Hindi", "English")
    out = content.decode("utf-8")
    assert re.search(r"[A-Za-z]{3,}", out) and not DEVANAGARI.search(out)


def test_hindi_to_tamil(client, require_model):
    require_model("indic-en")  # indic-indic model, or English pivot through both 200M models
    require_model("en-indic")
    _, content = translate_ok(client, "h.txt", "आज मौसम बहुत अच्छा है।".encode("utf-8"), "Hindi", "Tamil")
    assert TAMIL.search(content.decode("utf-8"))


def test_downloads_survive_restart(client, require_model):
    require_model("en-indic")
    result, content = translate_ok(client, "notes.txt", TXT_TEXT.encode("utf-8"))
    main.JOB_CACHE.clear()                                      # simulate a backend restart
    again = client.get(result["download_url"])
    assert again.status_code == 200 and again.content == content
    assert client.get(f"/api/jobs/{result['job_id']}").json()["status"] == "completed"


@pytest.mark.parametrize("path,source,target", [
    ("test_indic.pdf", "Hindi", "English"),
    ("backend/storage/uploads/one_page_english_conversation_7942985b.pdf", "English", "Hindi"),
])
def test_repository_sample_pdfs(client, require_model, path, source, target):
    sample = ASSETS_DIR / path
    if not sample.exists():
        pytest.skip(f"{path} not present")
    require_model("indic-en" if source != "English" else "en-indic")
    result, content = translate_ok(client, sample.name, sample.read_bytes(), source, target)
    out = pymupdf.open(stream=content, filetype="pdf")
    assert out.page_count == pymupdf.open(sample).page_count
    assert result["chunk_count"] > 0
