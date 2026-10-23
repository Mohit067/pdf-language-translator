"""Shared fixtures. Sample documents are generated as real files with the same libraries
users' documents come from; no translation output is faked anywhere in this suite."""

import io
import time
from pathlib import Path

import docx
import openpyxl
import pymupdf
import pytest
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fastapi.testclient import TestClient

import backend.document_processors as processors
import backend.main as main
from backend.agent_pipeline import ModelLoadError, agent

ASSETS_DIR = Path(__file__).resolve().parents[2]


@pytest.fixture(autouse=True)
def isolated_storage(tmp_path, monkeypatch):
    uploads, outputs = tmp_path / "uploads", tmp_path / "outputs"
    uploads.mkdir()
    outputs.mkdir()
    monkeypatch.setattr(main, "UPLOADS_DIR", uploads)
    monkeypatch.setattr(main, "OUTPUTS_DIR", outputs)
    monkeypatch.setattr(processors, "OUTPUTS_DIR", outputs)
    monkeypatch.setattr(main, "PRELOAD_MODELS", [])  # tests load the model only where they need it
    monkeypatch.setattr(main, "_preload_models", lambda: None)  # no startup network checks in tests
    main.JOB_CACHE.clear()
    yield uploads, outputs


@pytest.fixture
def client():
    with TestClient(main.app) as c:
        yield c


_model_state = {}


def _model_ready(direction: str) -> str:
    if direction not in _model_state:
        try:
            agent.load_model(direction)
            _model_state[direction] = ""
        except ModelLoadError as e:
            _model_state[direction] = str(e)
    return _model_state[direction]


@pytest.fixture
def require_model():
    def _require(direction: str = "en-indic"):
        reason = _model_ready(direction)
        if reason:
            pytest.skip(f"Local IndicTrans2 model unavailable: {reason[:200]}")
    return _require


def submit(client, filename: str, data: bytes, source="English", target="Hindi"):
    return client.post(
        "/api/translate",
        files={"file": (filename, data)},
        data={"source_lang": source, "target_lang": target, "tone": "Professional"},
    )


def wait_for_job(client, job_id: str, timeout: float = 900) -> dict:
    deadline = time.time() + timeout
    while time.time() < deadline:
        status = client.get(f"/api/jobs/{job_id}").json()
        if status["status"] in ("completed", "failed"):
            return status
        time.sleep(0.3)
    raise AssertionError(f"Job {job_id} did not finish within {timeout}s")


# ---------------------------------------------------------------------- document builders
def make_pdf(paragraphs=None, with_image=False) -> bytes:
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((200, 70), "Quarterly Business Report", fontsize=18, fontname="helv")
    page.draw_rect(pymupdf.Rect(50, 90, 545, 92), color=(0.2, 0.3, 0.6), fill=(0.2, 0.3, 0.6))
    paragraphs = paragraphs or [
        "The company expanded into three new regional markets this quarter.",
        "Customer satisfaction improved because support teams answered requests faster.",
    ]
    y = 130
    for text in paragraphs:
        page.insert_textbox(pymupdf.Rect(50, y, 545, y + 60), text, fontsize=11, fontname="helv")
        y += 70
    page.insert_text((50, y + 20), "Total amount", fontsize=11, fontname="helv")
    page.insert_text((450, y + 20), "4,500.00", fontsize=11, fontname="helv")  # separate column
    if with_image:
        pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 40, 40), 0)
        pix.set_rect(pix.irect, (200, 30, 30))
        page.insert_image(pymupdf.Rect(450, 700, 530, 780), pixmap=pix)
    return doc.tobytes()


def make_scanned_pdf() -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 50, 50), 0)
    pix.set_rect(pix.irect, (240, 240, 240))
    page.insert_image(page.rect, pixmap=pix)
    return doc.tobytes()


def make_encrypted_pdf() -> bytes:
    doc = pymupdf.open()
    doc.new_page().insert_text((50, 50), "Secret text", fontname="helv")
    return doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="userpw", owner_pw="ownerpw")


def make_docx() -> bytes:
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "Confidential company report"
    h = d.add_heading("Annual Summary", level=1)
    p = d.add_paragraph()
    p.add_run("The project was completed ").bold = True
    p.add_run("on time and within budget.")
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 20, 20), 0)
    pix.set_rect(pix.irect, (10, 120, 200))
    d.add_paragraph().add_run().add_picture(io.BytesIO(pix.tobytes("png")))
    t = d.add_table(rows=2, cols=3)
    t.cell(0, 0).merge(t.cell(0, 1)).text = "Customer feedback"
    t.cell(0, 2).text = "Status"
    t.cell(1, 0).text = "Delivery speed"
    t.cell(1, 1).text = "12,500"
    t.cell(1, 2).text = "Approved by the manager"
    c = d.add_paragraph("Thank you for your business.")
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Product", "Units", "Price", "Revenue"])
    ws.append(["Laptop bag", 10, 25.5, "=B2*C2"])
    ws.append(["Office chair", 4, 120, "=B3*C3"])
    ws.append(["Total", "=SUM(B2:B3)", None, "=SUM(D2:D3)"])
    ws["F1"] = "INV-2026-001"
    ws2 = wb.create_sheet("Notes")
    ws2["A1"] = "Laptop bag"
    ws2["A2"] = "Prices include tax."
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


CSV_TEXT = 'Name,Department,Comment\r\nAsha,Finance,"Approved the budget, with changes"\r\nRavi,Sales,Pending review\r\n101,202,303\r\n'

TXT_TEXT = "Meeting Notes\n\nThe team agreed to launch the product next month.\nMarketing will prepare the campaign.\n\nTotal: 4,500.00\n"

MD_TEXT = """# Project Overview

This tool translates documents.

- Fast processing
- Works offline

```python
print("do not translate")
```

| Feature | Status |
|---------|--------|
| Export  | Ready  |
"""
