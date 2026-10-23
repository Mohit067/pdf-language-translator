import codecs
import csv
import html
import io
import json
import logging
import re
import statistics
import uuid
import zipfile
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional, Tuple

import docx
import openpyxl
import pymupdf
from docx.oxml.ns import qn

from backend.config import (
    MAX_PDF_PAGES,
    OUTPUTS_DIR,
    RTL_LANG_CODES,
    get_lang_code,
)
from backend.agent_pipeline import ModelLoadError, agent, model_direction, needs_translation

logger = logging.getLogger("document_processor")
logging.basicConfig(level=logging.INFO)

# progress(stage 1-4, percent 0-100, message)
ProgressFn = Callable[[int, float, str], None]


class DocumentError(ValueError):
    """The uploaded document cannot be processed (corrupted, encrypted, empty, scanned...)."""


def _noop_progress(stage: int, percent: float, message: str) -> None:
    pass


class DocumentProcessor:
    """Handles parsing, segmenting, translating, and reconstructing entire documents."""

    @staticmethod
    def process_document(
        file_path: Path,
        source_lang: str,
        target_lang: str,
        tone: str = "Professional",
        job_id: Optional[str] = None,
        progress: Optional[ProgressFn] = None,
    ) -> Dict[str, Any]:
        ext = file_path.suffix.lower()
        job_id = job_id or str(uuid.uuid4())
        ctx = _Context(file_path, job_id, source_lang, target_lang, progress or _noop_progress)
        ctx.progress(1, 2, "Parsing document")

        if ext == ".pdf":
            return DocumentProcessor._process_pdf(ctx)
        elif ext in [".xlsx", ".xls"]:
            return DocumentProcessor._process_excel(ctx)
        elif ext == ".docx":
            return DocumentProcessor._process_docx(ctx)
        elif ext == ".csv":
            return DocumentProcessor._process_csv(ctx)
        elif ext in [".txt", ".md"]:
            return DocumentProcessor._process_text(ctx)
        else:
            raise DocumentError(f"Unsupported file extension: {ext}")

    # ------------------------------------------------------------------ PDF
    @staticmethod
    def _process_pdf(ctx: "_Context") -> Dict[str, Any]:
        try:
            doc = pymupdf.open(str(ctx.file_path), filetype="pdf")
        except Exception as e:
            raise DocumentError(f"The PDF could not be opened; it may be corrupted or not a real PDF ({e}).") from e

        with doc:
            if doc.needs_pass:
                raise DocumentError("The PDF is password-protected. Remove the password and upload it again.")
            if doc.page_count == 0:
                raise DocumentError("The PDF has no pages.")
            if doc.page_count > MAX_PDF_PAGES:
                raise DocumentError(f"The PDF has {doc.page_count} pages; the limit is {MAX_PDF_PAGES} (MAX_PDF_PAGES).")

            units: List[_PdfUnit] = []
            for page in doc:
                try:
                    units.extend(_extract_pdf_units(page))
                except Exception as e:
                    logger.warning(f"Text extraction failed on page {page.number + 1}: {e}")
                    ctx.warnings.append(f"Page {page.number + 1}: text could not be extracted and was left unchanged.")

            to_translate = [u for u in units if needs_translation(u.text)]
            if not to_translate:
                raise DocumentError(
                    "No selectable text was found in this PDF. It is probably a scanned image; "
                    "OCR is not supported, so please upload a PDF with a text layer."
                )
            textless = doc.page_count - len({u.page for u in to_translate})
            if textless:
                ctx.warnings.append(f"{textless} page(s) contain no selectable text (scanned images are not OCR'd) and were kept as-is.")

            translations = ctx.translate([u.text for u in to_translate])

            ctx.progress(4, 92, "Rebuilding PDF with original layout")
            rtl = ctx.tgt_code in RTL_LANG_CODES
            by_page: Dict[int, List[Tuple[_PdfUnit, str]]] = {}
            for unit, translated in zip(to_translate, translations):
                by_page.setdefault(unit.page, []).append((unit, translated))

            for pno, items in by_page.items():
                page = doc[pno]
                for unit, _ in items:
                    for rect in unit.line_rects:
                        page.add_redact_annot(rect, fill=False)
                page.apply_redactions(
                    images=pymupdf.PDF_REDACT_IMAGE_NONE,
                    graphics=pymupdf.PDF_REDACT_LINE_ART_NONE,
                    text=pymupdf.PDF_REDACT_TEXT_REMOVE,
                )
                for unit, translated in items:
                    _render_pdf_unit(page, unit, translated, rtl)

            try:
                doc.subset_fonts()
            except Exception as e:
                logger.debug(f"Font subsetting skipped: {e}")

            output_filename = f"translated_{ctx.job_id}.pdf"
            output_path = OUTPUTS_DIR / output_filename
            doc.save(str(output_path), garbage=3, deflate=True)
            num_pages = doc.page_count

        return ctx.result(
            output_filename, output_path, "pdf", num_pages,
            [u.text for u in to_translate], translations, [u.page + 1 for u in to_translate],
        )

    # ------------------------------------------------------------------ Excel
    @staticmethod
    def _process_excel(ctx: "_Context") -> Dict[str, Any]:
        if ctx.file_path.suffix.lower() == ".xls":
            wb = _load_xls_as_workbook(ctx.file_path)
            ctx.warnings.append("Legacy .xls was converted to .xlsx: cell values are translated, but formulas and formatting from the .xls are not carried over.")
        else:
            try:
                wb = openpyxl.load_workbook(str(ctx.file_path))
            except Exception as e:
                raise DocumentError(f"The Excel file could not be opened; it may be corrupted or password-protected ({e}).") from e

        cells: List[Tuple[int, Any]] = []
        for sheet_idx, ws in enumerate(wb.worksheets):
            for row in ws.iter_rows():
                for cell in row:
                    val = cell.value
                    # data_type "s" excludes formulas ("f"), numbers, dates and booleans
                    if cell.data_type == "s" and isinstance(val, str) and needs_translation(val):
                        cells.append((sheet_idx, cell))

        if not cells:
            raise DocumentError("No translatable text was found in the workbook (only numbers, dates or formulas).")

        first_sheet: Dict[str, int] = {}
        for sheet_idx, c in cells:
            first_sheet.setdefault(c.value.strip(), sheet_idx + 1)
        unique = list(first_sheet)
        translated_unique = dict(zip(unique, ctx.translate(unique)))

        ctx.progress(4, 92, "Writing translated workbook")
        for _, cell in cells:
            cell.value = translated_unique[cell.value.strip()]

        output_filename = f"translated_{ctx.job_id}.xlsx"
        output_path = OUTPUTS_DIR / output_filename
        wb.save(str(output_path))

        return ctx.result(
            output_filename, output_path, "xlsx", len(wb.worksheets),
            unique, [translated_unique[u] for u in unique], [first_sheet[u] for u in unique],
        )

    # ------------------------------------------------------------------ Word
    @staticmethod
    def _process_docx(ctx: "_Context") -> Dict[str, Any]:
        try:
            doc = docx.Document(str(ctx.file_path))
        except Exception as e:
            raise DocumentError(f"The Word file could not be opened; it may be corrupted or an old .doc renamed to .docx ({e}).") from e

        paragraphs = []
        for p_el in _iter_docx_paragraphs(doc):
            t_elements = _own_text_elements(p_el)
            text = "".join(t.text or "" for t in t_elements)
            if needs_translation(text):
                paragraphs.append((t_elements, text.strip()))

        if not paragraphs:
            raise DocumentError("No translatable text was found in the Word document.")

        translations = ctx.translate([text for _, text in paragraphs])

        ctx.progress(4, 92, "Writing translated Word document")
        for (t_elements, _), translated in zip(paragraphs, translations):
            # Keep the first run's formatting and every non-text run (images, fields, breaks)
            t_elements[0].text = translated
            t_elements[0].set(qn("xml:space"), "preserve")
            for t in t_elements[1:]:
                t.text = ""

        output_filename = f"translated_{ctx.job_id}.docx"
        output_path = OUTPUTS_DIR / output_filename
        doc.save(str(output_path))

        return ctx.result(
            output_filename, output_path, "docx", 1,
            [text for _, text in paragraphs], translations, [1] * len(paragraphs),
        )

    # ------------------------------------------------------------------ CSV
    @staticmethod
    def _process_csv(ctx: "_Context") -> Dict[str, Any]:
        content, _, newline = _read_text_file(ctx.file_path)
        try:
            dialect = csv.Sniffer().sniff(content[:65536], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(io.StringIO(content), dialect))

        unique = list(dict.fromkeys(cell.strip() for row in rows for cell in row if needs_translation(cell)))
        if not unique:
            raise DocumentError("No translatable text was found in the CSV file.")
        translated_unique = dict(zip(unique, ctx.translate(unique)))

        ctx.progress(4, 92, "Writing translated CSV")
        out_rows = [[translated_unique.get(cell.strip(), cell) for cell in row] for row in rows]
        buffer = io.StringIO()
        writer = csv.writer(buffer, delimiter=dialect.delimiter, quotechar=dialect.quotechar or '"',
                            quoting=csv.QUOTE_MINIMAL, lineterminator=newline)
        writer.writerows(out_rows)

        output_filename = f"translated_{ctx.job_id}.csv"
        output_path = OUTPUTS_DIR / output_filename
        # BOM so Excel detects UTF-8 and renders Indic scripts correctly
        output_path.write_text(buffer.getvalue(), encoding="utf-8-sig", newline="")

        changed = [i for i, row in enumerate(rows) if any(needs_translation(c) for c in row)]
        return ctx.result(
            output_filename, output_path, "csv", 1,
            [dialect.delimiter.join(rows[i]) for i in changed],
            [dialect.delimiter.join(out_rows[i]) for i in changed],
            [1] * len(changed),
        )

    # ------------------------------------------------------------------ TXT / Markdown
    @staticmethod
    def _process_text(ctx: "_Context") -> Dict[str, Any]:
        content, bom, newline = _read_text_file(ctx.file_path)
        is_md = ctx.file_path.suffix.lower() == ".md"
        lines = content.split("\n")

        # Each translatable unit: (line index, prefix, text, suffix)
        units: List[Tuple[int, str, str, str]] = []
        in_fence = False
        for i, line in enumerate(lines):
            if is_md:
                if re.match(r"^\s*(```|~~~)", line):
                    in_fence = not in_fence
                    continue
                if in_fence or re.match(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?\s*$", line) or re.match(r"^\s*<", line):
                    continue
                if line.lstrip().startswith("|"):
                    units.append((i, "", line, ""))  # table row, cells handled below
                    continue
                m = re.match(r"^(\s*(?:>\s*)*(?:#{1,6}\s+|[-*+]\s+(?:\[[ xX]\]\s+)?|\d+[.)]\s+)?)(.*?)(\s*)$", line)
                prefix, body, suffix = m.group(1), m.group(2), m.group(3)
            else:
                prefix, body, suffix = "", line, ""
            if needs_translation(body):
                units.append((i, prefix, body, suffix))

        # Markdown table rows translate cell-by-cell to keep the pipes intact
        texts: List[str] = []
        layout: List[Tuple[int, str, Optional[List[str]], str]] = []
        for i, prefix, body, suffix in units:
            if is_md and body.lstrip().startswith("|"):
                cells = body.split("|")
                layout.append((i, prefix, cells, suffix))
                texts.extend(c.strip() for c in cells if needs_translation(c))
            else:
                layout.append((i, prefix, None, suffix))
                texts.append(body)

        if not texts:
            raise DocumentError("No translatable text was found in the file.")

        translated_texts = iter(ctx.translate(texts))
        out_lines = list(lines)
        for i, prefix, cells, suffix in layout:
            if cells is None:
                out_lines[i] = f"{prefix}{next(translated_texts)}{suffix}"
            else:
                out_lines[i] = "|".join(
                    f" {next(translated_texts)} " if needs_translation(c) else c for c in cells
                )

        ctx.progress(4, 92, "Writing translated file")
        translated_content = "\n".join(out_lines)
        output_filename = f"translated_{ctx.job_id}{ctx.file_path.suffix.lower()}"
        output_path = OUTPUTS_DIR / output_filename
        output_path.write_text(translated_content.replace("\n", newline), encoding="utf-8-sig" if bom else "utf-8", newline="")

        # Paragraph-aligned chunks for preview (line structure is preserved 1:1)
        originals, translations = [], []
        for start, end in _paragraph_spans(lines):
            orig = "\n".join(lines[start:end])
            if needs_translation(orig):
                originals.append(orig)
                translations.append("\n".join(out_lines[start:end]))
        return ctx.result(
            output_filename, output_path, ctx.file_path.suffix.lower().lstrip("."), 1,
            originals, translations, [1] * len(originals),
        )


# ====================================================================== helpers
class _Context:
    def __init__(self, file_path: Path, job_id: str, source_lang: str, target_lang: str, progress: ProgressFn):
        self.file_path = file_path
        self.job_id = job_id
        self.source_lang = source_lang
        self.target_lang = target_lang
        self.src_code = get_lang_code(source_lang)
        self.tgt_code = get_lang_code(target_lang)
        if self.src_code == self.tgt_code:
            raise ValueError("Source and target languages must be different.")
        self.progress = progress
        self.warnings: List[str] = []

    def translate(self, texts: List[str]) -> List[str]:
        direction = model_direction(self.src_code, self.tgt_code)
        self.progress(2, 6, "Loading local IndicTrans2 model (first run downloads the weights)")
        try:
            agent.load_model(direction)
        except ModelLoadError:
            if direction != "indic-indic":  # indic-indic can pivot through English
                raise

        def on_progress(done: int, total: int) -> None:
            pct = 10 + 80 * (done / total if total else 1)
            self.progress(3, pct, f"Translating segments ({done}/{total})")

        # Checkpoint: every finished batch is appended to disk, so a job resumed after a restart
        # (see main.resume_unfinished_jobs) only translates what is still missing.
        known = self._load_checkpoint()
        if known:
            logger.info(f"Job {self.job_id}: resuming with {len(known)} already-translated sentences")
        with open(self.checkpoint_path, "a", encoding="utf-8") as checkpoint:
            def on_batch(pairs: Dict[str, str]) -> None:
                for src, out in pairs.items():
                    checkpoint.write(json.dumps({"s": src, "t": out}, ensure_ascii=False) + "\n")
                checkpoint.flush()

            return agent.translate_texts(texts, self.source_lang, self.target_lang, on_progress,
                                         known=known, on_batch=on_batch)

    @property
    def checkpoint_path(self) -> Path:
        return OUTPUTS_DIR / f"checkpoint_{self.job_id}.jsonl"

    def _load_checkpoint(self) -> Dict[str, str]:
        known: Dict[str, str] = {}
        if self.checkpoint_path.exists():
            with open(self.checkpoint_path, encoding="utf-8") as f:
                for line in f:
                    try:
                        item = json.loads(line)
                        known[item["s"]] = item["t"]
                    except (ValueError, KeyError):
                        continue  # a line cut off by the interruption
        return known

    def result(self, output_filename: str, output_path: Path, output_format: str, pages: int,
               originals: List[str], translations: List[str], page_numbers: List[int]) -> Dict[str, Any]:
        chunks = [
            {
                "chunk_id": i + 1,
                "page": page_numbers[i],
                "original": orig,
                "translated": trans,
                "word_count_original": len(orig.split()),
                "word_count_translated": len(trans.split()),
                "src_code": self.src_code,
                "tgt_code": self.tgt_code,
            }
            for i, (orig, trans) in enumerate(zip(originals, translations))
        ]
        self.progress(4, 100, "Done")
        self.checkpoint_path.unlink(missing_ok=True)
        return {
            "job_id": self.job_id,
            "filename": self.file_path.name,
            "output_filename": output_filename,
            "output_format": output_format,
            "original_path": str(self.file_path),
            "output_path": str(output_path),
            "pages": pages,
            "chunk_count": len(chunks),
            "source_lang": self.source_lang,
            "target_lang": self.target_lang,
            "chunks": chunks,
            "full_original": "\n\n".join(originals),
            "full_translated": "\n\n".join(translations),
            "warnings": self.warnings,
        }


# ---------------------------------------------------------------------- PDF helpers
class _PdfUnit:
    """A paragraph (or isolated line fragment) of a PDF page with its geometry and style."""

    def __init__(self, page: int):
        self.page = page
        self.lines: List[Dict[str, Any]] = []   # {"text", "rect", "size", "bold", ...}

    @property
    def line_rects(self) -> List[pymupdf.Rect]:
        return [l["rect"] for l in self.lines]

    @property
    def rect(self) -> pymupdf.Rect:
        r = pymupdf.Rect(self.lines[0]["rect"])
        for l in self.lines[1:]:
            r |= l["rect"]
        return r

    @property
    def text(self) -> str:
        if len(self.lines) == 1:
            return self.lines[0]["text"].strip()
        right = max(l["rect"].x1 for l in self.lines)
        width = self.rect.width or 1
        parts = []
        for i, l in enumerate(self.lines):
            parts.append(l["text"].strip())
            if i < len(self.lines) - 1:
                # A line ending well before the paragraph's right edge is a deliberate line break
                parts.append("\n" if right - l["rect"].x1 > 0.2 * width else " ")
        return "".join(parts).strip()

    def style(self) -> Dict[str, Any]:
        sizes = [l["size"] for l in self.lines]
        first = self.lines[0]
        return {"size": statistics.median(sizes), "color": first["color"], "bold": first["bold"],
                "italic": first["italic"], "serif": first["serif"]}


def _span_style(span: Dict[str, Any]) -> Dict[str, Any]:
    font = span.get("font", "").lower()
    flags = span.get("flags", 0)
    return {
        "size": round(span.get("size", 10), 1),
        "color": span.get("color", 0),
        "bold": bool(flags & 16) or "bold" in font or "black" in font,
        "italic": bool(flags & 2) or "italic" in font or "oblique" in font,
        "serif": bool(flags & 4) or any(k in font for k in ("times", "serif", "georgia", "garamond", "cambria")),
    }


def _extract_pdf_units(page: pymupdf.Page) -> List[_PdfUnit]:
    data = page.get_text("dict", flags=pymupdf.TEXTFLAGS_TEXT | pymupdf.TEXT_DEHYPHENATE)
    units: List[_PdfUnit] = []
    for block in data.get("blocks", []):
        if block.get("type") != 0:
            continue
        current: Optional[_PdfUnit] = None
        for line in block.get("lines", []):
            dx, dy = line.get("dir", (1, 0))
            if abs(dy) > 0.01 or dx < 0:   # rotated/vertical text is left untouched
                current = None
                continue
            # Split a line into fragments where spans are far apart (table columns, tab stops)
            fragments: List[List[Dict[str, Any]]] = []
            for span in line.get("spans", []):
                if not span.get("text"):
                    continue
                if fragments:
                    prev = fragments[-1][-1]
                    gap = span["bbox"][0] - prev["bbox"][2]
                    if gap > max(2.5 * span.get("size", 10), 12):
                        fragments.append([span])
                        continue
                    fragments[-1].append(span)
                else:
                    fragments.append([span])
            fragments = [f for f in fragments if "".join(s["text"] for s in f).strip()]
            if not fragments:
                continue

            entries = []
            for frag in fragments:
                rect = pymupdf.Rect(frag[0]["bbox"])
                for s in frag[1:]:
                    rect |= pymupdf.Rect(s["bbox"])
                main = max(frag, key=lambda s: len(s["text"].strip()))
                entries.append({"text": "".join(s["text"] for s in frag), "rect": rect, **_span_style(main)})

            if len(entries) > 1:
                current = None
                for e in entries:
                    u = _PdfUnit(page.number)
                    u.lines.append(e)
                    units.append(u)
                continue

            entry = entries[0]
            if current is not None:
                last = current.lines[-1]
                same_style = (abs(last["size"] - entry["size"]) < 1.0 and last["bold"] == entry["bold"])
                close = entry["rect"].y0 - last["rect"].y1 < 0.8 * entry["size"]
                overlaps = entry["rect"].x0 < current.rect.x1 and entry["rect"].x1 > current.rect.x0
                if same_style and close and overlaps:
                    current.lines.append(entry)
                    continue
            current = _PdfUnit(page.number)
            current.lines.append(entry)
            units.append(current)
    return units


def _render_pdf_unit(page: pymupdf.Page, unit: _PdfUnit, translated: str, rtl: bool) -> None:
    style = unit.style()
    rect = pymupdf.Rect(unit.rect)
    page_w = page.rect.width
    center = (rect.x0 + rect.x1) / 2
    centered = (len(unit.lines) == 1 and rect.width < 0.75 * page_w and abs(center - page_w / 2) < 0.04 * page_w)
    if centered:
        # Give centered headings room to grow symmetrically instead of shrinking the font
        half = max(rect.width / 2, min(center - 36, page_w - 36 - center))
        rect.x0, rect.x1 = center - half, center + half
    align = "center" if centered else ("right" if rtl else "left")
    css = (
        "* {margin: 0; padding: 0;} "
        f"body {{font-family: {'serif' if style['serif'] else 'sans-serif'}; font-size: {style['size']}px; "
        f"line-height: 1.15; color: #{style['color']:06x}; text-align: {align}; "
        f"font-weight: {'bold' if style['bold'] else 'normal'}; font-style: {'italic' if style['italic'] else 'normal'};"
        f"{' direction: rtl;' if rtl else ''}}}"
    )
    body = html.escape(translated).replace("\n", "<br>")
    try:
        page.insert_htmlbox(rect, body, css=css, scale_low=0)
    except Exception as e:
        logger.warning(f"Could not place translated text on page {page.number + 1}: {e}")


# ---------------------------------------------------------------------- DOCX helpers
def _iter_docx_paragraphs(doc):
    """Yields every w:p element once: body (incl. nested tables and content controls),
    headers/footers, and text boxes. Merged cells are unique <w:tc> elements in the XML,
    so iterating XML (instead of row.cells) never visits a cell twice."""
    def walk(container):
        for child in container.iterchildren():
            if child.tag == qn("w:p"):
                yield child
            elif child.tag == qn("w:tbl"):
                for tr in child.iterchildren(qn("w:tr")):
                    for tc in tr.iterchildren(qn("w:tc")):
                        yield from walk(tc)
            elif child.tag in (qn("w:sdt"), qn("w:customXml")):
                content = child.find(qn("w:sdtContent"))
                yield from walk(content if content is not None else child)

    roots = [doc.element.body]
    seen_parts = set()
    for section in doc.sections:
        for hf in (section.header, section.first_page_header, section.even_page_header,
                   section.footer, section.first_page_footer, section.even_page_footer):
            if hf.is_linked_to_previous:
                continue  # accessing a linked header's element would create a new definition
            part_id = id(hf.part)
            if part_id not in seen_parts:
                seen_parts.add(part_id)
                roots.append(hf._element)

    for root in roots:
        yield from walk(root)
        for txbx in root.iter(qn("w:txbxContent")):
            yield from walk(txbx)


def _own_text_elements(p_el) -> List[Any]:
    """w:t elements belonging to this paragraph (excluding nested text-box paragraphs)."""
    own = []
    for t in p_el.iter(qn("w:t")):
        parent = t.getparent()
        while parent is not None and parent.tag != qn("w:p"):
            parent = parent.getparent()
        if parent is p_el:
            own.append(t)
    return own


# ---------------------------------------------------------------------- Excel helpers
def _load_xls_as_workbook(path: Path) -> "openpyxl.Workbook":
    try:
        import xlrd
    except ImportError as e:
        raise DocumentError("Reading legacy .xls files requires the 'xlrd' package; save the file as .xlsx instead.") from e
    try:
        book = xlrd.open_workbook(str(path))
    except Exception as e:
        raise DocumentError(f"The .xls file could not be opened; it may be corrupted ({e}).") from e

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet in book.sheets():
        ws = wb.create_sheet(title=sheet.name[:31] or "Sheet")
        for r in range(sheet.nrows):
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                    continue
                value = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    value = xlrd.xldate_as_datetime(value, book.datemode)
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    value = bool(value)
                elif cell.ctype == xlrd.XL_CELL_NUMBER and float(value).is_integer():
                    value = int(value)
                ws.cell(row=r + 1, column=c + 1, value=value)
        for rlo, rhi, clo, chi in sheet.merged_cells:
            ws.merge_cells(start_row=rlo + 1, end_row=rhi, start_column=clo + 1, end_column=chi)
    return wb


# ---------------------------------------------------------------------- text helpers
def _read_text_file(path: Path) -> Tuple[str, bool, str]:
    """Returns (content with \\n newlines, had_utf8_bom, original newline sequence)."""
    raw = path.read_bytes()
    bom = False
    if raw.startswith(codecs.BOM_UTF8):
        content, bom = raw[len(codecs.BOM_UTF8):].decode("utf-8", errors="replace"), True
    elif raw.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        content = raw.decode("utf-16")
    else:
        if b"\x00" in raw[:8192]:
            raise DocumentError("The file appears to be binary, not plain text.")
        try:
            content = raw.decode("utf-8")
        except UnicodeDecodeError:
            content = raw.decode("cp1252", errors="replace")
            logger.info(f"{path.name} is not UTF-8; decoded as Windows-1252")
    newline = "\r\n" if "\r\n" in content else "\n"
    return content.replace("\r\n", "\n").replace("\r", "\n"), bom, newline


def _paragraph_spans(lines: List[str]) -> List[Tuple[int, int]]:
    spans, start = [], None
    for i, line in enumerate(lines):
        if line.strip():
            if start is None:
                start = i
        elif start is not None:
            spans.append((start, i))
            start = None
    if start is not None:
        spans.append((start, len(lines)))
    return spans
